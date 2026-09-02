import csv
import json
import os
import re
from datetime import date, datetime


class Sheet:
    """Minimal spreadsheet view shared by .xls, .xlsx, and .csv."""

    def __init__(self, rows):
        self._rows = rows
        self.nrows = len(rows)
        self.ncols = max((len(row) for row in rows), default=0)

    def cell_value(self, row_idx, col_idx):
        if row_idx < 0 or row_idx >= self.nrows:
            return ""
        row = self._rows[row_idx]
        if col_idx < 0 or col_idx >= len(row):
            return ""
        return row[col_idx]


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def extract_city_state_zip_from_line(text):
    text = clean_text(text)
    match = re.search(r"^(.*?),\s*([A-Z]{2})\s+(\d{5})(?:-\d{4})?$", text)
    if match:
        city = match.group(1).strip()
        state = match.group(2).strip()
        zip_code = match.group(3).strip()
        return city, state, zip_code
    return "", "", ""


def find_cell_by_text(sheet, target_text):
    target_text = target_text.strip().lower()

    for row_idx in range(sheet.nrows):
        for col_idx in range(sheet.ncols):
            value = clean_text(sheet.cell_value(row_idx, col_idx)).lower()
            if value == target_text:
                return row_idx, col_idx
    return None, None


def to_int_string(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)

    text = clean_text(value)
    if text.endswith(".0"):
        return text[:-2]
    return text


def _normalize_rows(raw_rows):
    rows = []
    for raw_row in raw_rows:
        row = list(raw_row)
        while row and clean_text(row[-1]) == "":
            row.pop()
        rows.append(row)
    while rows and not any(clean_text(cell) for cell in rows[-1]):
        rows.pop()
    return rows


def _load_xls(file_path, sheet_index=0):
    try:
        import xlrd
    except ImportError as exc:
        raise ImportError("xlrd is required to parse .xls files") from exc

    workbook = xlrd.open_workbook(file_path)
    worksheet = workbook.sheet_by_index(sheet_index)
    rows = [
        [worksheet.cell_value(row_idx, col_idx) for col_idx in range(worksheet.ncols)]
        for row_idx in range(worksheet.nrows)
    ]
    return Sheet(_normalize_rows(rows))


def _load_xlsx(file_path, sheet_index=0):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is required to parse .xlsx files") from exc

    workbook = load_workbook(file_path, data_only=True, read_only=True)
    try:
        worksheet = workbook.worksheets[sheet_index]
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            rows.append(list(row))
        return Sheet(_normalize_rows(rows))
    finally:
        workbook.close()


def _load_csv(file_path, sheet_index=0):
    if sheet_index != 0:
        raise ValueError("CSV files have a single sheet; sheet_index must be 0")

    with open(file_path, "r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [list(row) for row in csv.reader(handle, dialect)]
    return Sheet(_normalize_rows(rows))


def load_sheet(file_path, sheet_index=0):
    extension = os.path.splitext(file_path)[1].lower()
    loaders = {
        ".xls": _load_xls,
        ".xlsx": _load_xlsx,
        ".xlsm": _load_xlsx,
        ".csv": _load_csv,
    }
    loader = loaders.get(extension)
    if loader is None:
        raise ValueError(
            f"Unsupported file type '{extension}'. Expected .xls, .xlsx, .xlsm, or .csv"
        )
    return loader(file_path, sheet_index)


def _value_after_label(cell_text, label):
    text = clean_text(cell_text)
    label_text = label.strip()
    if text.lower().startswith(label_text.lower()):
        remainder = text[len(label_text):].strip().lstrip(":").strip()
        return remainder
    return ""


def get_labeled_value(sheet, label):
    row, col = find_cell_by_text(sheet, label)
    if row is not None:
        if col + 1 < sheet.ncols:
            adjacent = clean_text(sheet.cell_value(row, col + 1))
            if adjacent:
                return adjacent
        same_cell = _value_after_label(sheet.cell_value(row, col), label)
        if same_cell:
            return same_cell

    label_lower = label.strip().lower()
    for row_idx in range(sheet.nrows):
        for col_idx in range(sheet.ncols):
            value = clean_text(sheet.cell_value(row_idx, col_idx))
            if not value.lower().startswith(label_lower):
                continue
            remainder = _value_after_label(value, label)
            if remainder:
                return remainder
            if col_idx + 1 < sheet.ncols:
                adjacent = clean_text(sheet.cell_value(row_idx, col_idx + 1))
                if adjacent:
                    return adjacent
    return ""


def extract_items(sheet):
    order_details = []

    qty_row, qty_col = find_cell_by_text(sheet, "Qty Ordered")
    part_row, part_col = find_cell_by_text(sheet, "Part Number")

    if qty_row is None or part_row is None or qty_row != part_row:
        return order_details

    header_row = qty_row

    for row_idx in range(header_row + 1, sheet.nrows):
        qty_value = sheet.cell_value(row_idx, qty_col) if qty_col < sheet.ncols else ""
        part_value = sheet.cell_value(row_idx, part_col) if part_col < sheet.ncols else ""

        qty = to_int_string(qty_value)

        sku = clean_text(part_value)
        if sku.endswith(".0"):
            sku = sku[:-2]
        if not qty and not sku:
            continue

        if qty and not sku:
            continue

        if sku.lower().startswith("sub total"):
            continue

        if not re.fullmatch(r"\d+(?:\.\d+)?", qty):
            continue

        order_details.append({
            "quantity": qty,
            "sku": sku
        })

    return order_details


def extract_po_data(file_path, sheet_index=0):
    sheet = load_sheet(file_path, sheet_index=sheet_index)

    result = {
        "order_number": "",
        "customer_name": "",
        "address1": "",
        "city": "",
        "state": "",
        "zip": "",
        "country": "United States",
        "order_lines": []
    }

    result["order_number"] = get_labeled_value(sheet, "PO Number:")

    ship_row, ship_col = find_cell_by_text(sheet, "Ship To")
    if ship_row is not None:
        if ship_row + 1 < sheet.nrows:
            result["customer_name"] = clean_text(sheet.cell_value(ship_row + 1, ship_col))

        if ship_row + 3 < sheet.nrows:
            result["address1"] = clean_text(sheet.cell_value(ship_row + 3, ship_col))

        if ship_row + 4 < sheet.nrows:
            city_state_zip_line = clean_text(sheet.cell_value(ship_row + 4, ship_col))
            city, state, zip_code = extract_city_state_zip_from_line(city_state_zip_line)
            result["city"] = city
            result["state"] = state
            result["zip"] = zip_code

    result["order_lines"] = extract_items(sheet)

    return result


if __name__ == "__main__":
    import sys

    file_path = sys.argv[1] if len(sys.argv) > 1 else "153-030326.xls"
    data = extract_po_data(file_path)
    print(json.dumps(data, indent=2))
